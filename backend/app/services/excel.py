"""Excel helpers for the bulk imports: safe reading, friendly templates."""
from __future__ import annotations

import io
import re
import zipfile
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from ..errors import ApiError

MAX_FILE_BYTES = 2 * 1024 * 1024
MAX_UNZIPPED_BYTES = 30 * 1024 * 1024
MAX_ROWS = 1000


def _bad(message: str) -> ApiError:
    return ApiError(422, "BAD_FILE", message, extra={"fields": {"file": message}})


def norm_header(value: Any) -> str:
    return re.sub(r"[^a-z]", "", str(value or "").lower())


def read_rows(raw: bytes, columns: dict[str, str], required: list[str]) -> list[tuple[int, dict[str, str]]]:
    """Reads the first sheet. `columns` maps a canonical key -> the start of its normalised header text.

    Returns [(excel_row_number, {key: text})] without the empty rows."""
    if not raw:
        raise _bad("The file is empty.")
    if len(raw) > MAX_FILE_BYTES:
        raise _bad("The file is too large (2 MB maximum).")
    if not zipfile.is_zipfile(io.BytesIO(raw)):
        raise _bad("This is not an Excel (.xlsx) file. Please download the sample template and fill it in.")
    with zipfile.ZipFile(io.BytesIO(raw)) as z:  # protects against "zip bombs"
        infos = z.infolist()
        if len(infos) > 300 or sum(i.file_size for i in infos) > MAX_UNZIPPED_BYTES:
            raise _bad("This file looks unusual and was refused. Please use the sample template.")
    try:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception:
        raise _bad("The Excel file could not be read. Please save it as .xlsx and try again.") from None
    try:
        ws = wb.worksheets[0]
        iterator = ws.iter_rows(values_only=True)
        header = next(iterator, None)
        if not header:
            raise _bad("The first sheet is empty.")
        index: dict[str, int] = {}
        for pos, cell in enumerate(header):
            h = norm_header(cell)
            if not h:
                continue
            for key, prefix in columns.items():
                if key not in index and h.startswith(prefix):
                    index[key] = pos
        missing = [k for k in required if k not in index]
        if missing:
            raise _bad("These columns are missing in the first row: " + ", ".join(missing) + ". Please use the sample template.")
        rows: list[tuple[int, dict[str, str]]] = []
        for number, values in enumerate(iterator, start=2):
            if number > MAX_ROWS + 1:
                raise _bad(f"Too many rows (maximum {MAX_ROWS}).")
            item = {key: _cell(values[pos]) if pos < len(values) else "" for key, pos in index.items()}
            if any(item.values()):
                rows.append((number, item))
        return rows
    finally:
        wb.close()


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", str(value)).strip()


def yes_no(value: str) -> bool | None:
    v = value.strip().lower()
    if v == "":
        return False
    if v in ("yes", "y", "true", "1"):
        return True
    if v in ("no", "n", "false", "0"):
        return False
    return None


# ---------------------------------------------------------------------- templates
_HEAD_FILL = PatternFill("solid", start_color="4F46E5")
_HEAD_FONT = Font(bold=True, color="FFFFFF")
_NOTE_FONT = Font(italic=True, color="555555")


def build_template(title: str, headers: list[str], widths: list[int], examples: list[list[Any]], notes: list[str],
                   yes_no_columns: list[int] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(headers)
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
        cell = ws.cell(row=1, column=col)
        cell.fill, cell.font = _HEAD_FILL, _HEAD_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 34
    for example in examples:
        ws.append(example)
    ws.freeze_panes = "A2"
    if yes_no_columns:
        dv = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
        ws.add_data_validation(dv)
        for col in yes_no_columns:
            letter = get_column_letter(col)
            dv.add(f"{letter}2:{letter}1001")
    info = wb.create_sheet("How to fill")
    info.column_dimensions["A"].width = 110
    for line in notes:
        info.append([line])
    info["A1"].font = Font(bold=True, size=13)
    for row in info.iter_rows(min_row=2):
        for cell in row:
            cell.font = _NOTE_FONT
            cell.alignment = Alignment(wrap_text=True)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
