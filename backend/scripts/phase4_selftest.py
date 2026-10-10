"""Phase 4 self-test on the REAL seeded School of IT data (in-process API as the real admin; teachers are impersonated for their own screens).

  python -m scripts.phase4_selftest

  1. weekly-lectures override: the three over-full B.Sc. classes are made to fit, generation reaches 100 %
  2. publish: per class + all at once, versions, audit, "unpublished changes", re-publish notifies ONLY the teachers whose lectures changed
  3. what teachers see: today / week, combined sessions, co-taught lectures, holidays - and never a draft
  4. PDF + Excel exports (class, teacher, school master)
  5. notifications: the bell (live hub), mark read, one per person per event, and a REAL Web Push delivery to a local "browser" (decrypted)
  6. semester hours: a test holiday -> "short by X" -> suggested slots are clash-free -> approved -> notified
  7. cleanup: every draft, run, publication, notification, extra lecture, test holiday and override goes away; the seed is verified again
"""
from __future__ import annotations

import os

os.environ.setdefault("READ_CACHE_SECONDS", "0")   # these scripts also write straight to the database: always read fresh

import asyncio
import base64
import io
import json
import os
import sys
import threading
from datetime import date, datetime, timedelta, timezone
from http.server import BaseHTTPRequestHandler, HTTPServer
from time import monotonic

import httpx

from app import db
from app.config import get_settings
from app.main import app
from app.routers import drafts as drafts_router
from app.security import deps
from app.services import generator as gen
from app.services import notify
from app.services import terms as term_svc
from app.services import views as vw

from .phase3_selftest import wait_run
from .seed_college import Api, SeedError, Who, _fake_authenticate, load_user, log, verify

OK = BAD = 0
PUSH_ARRIVED: list[dict] = []


def check(cond: bool, msg: str, extra: str = "") -> None:
    global OK, BAD
    OK += bool(cond)
    BAD += not cond
    log(("  PASS  " if cond else "  FAIL  ") + msg + (f"   {extra}" if extra else ""))


class Acting:
    """Run a block of calls as another person (their own screens), then go back to the admin."""

    def __init__(self, name: str | None = None, role: str | None = None) -> None:
        self.name, self.role = name, role

    async def __aenter__(self):
        self.previous = Who.user
        email = await db.fetchval("select email from public.profiles where full_name = $1", self.name) if self.name else None
        Who.user = await load_user(email=email, role=self.role)
        return Who.user

    async def __aexit__(self, *exc):
        Who.user = self.previous


async def main() -> int:
    settings = get_settings()
    started = datetime.now(timezone.utc)
    async with app.router.lifespan_context(app):
        Who.user = await load_user(email=settings.seed_admin_email, role="admin")
        deps.authenticate = _fake_authenticate
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://localhost",
                                     headers={"Origin": "http://localhost:5173"}, timeout=300) as client:
            api = Api(client)
            original = {str(r["id"]): r["weekly_lectures"] for r in await db.fetch("select id, weekly_lectures from public.subjects")}
            try:
                await scenarios(api, client)
            except (SeedError, AssertionError) as exc:
                log(f"\nSTOPPED: {exc}")
                return 1
            finally:
                await notify.drain()
                await cleanup(api, started, original)
            log("\n== the seeded data is untouched ==")
            bad = await verify(api)
            check(bad == 0, "seed verifier: ALL PASS")
    log(f"\nPHASE 4 SELF-TEST: {OK} checks passed, {BAD} failed")
    return 1 if BAD else 0


# ====================================================================================================================
async def class_numbers(sid: str) -> dict[str, tuple[int, int, str]]:
    """label -> (periods needed a week, periods free, id) straight from the generation engine's own problem statement."""
    term = await term_svc.require_active_term()
    L = await drafts_router.load(__import__("uuid").UUID(sid), term)
    problem = gen.build_problem(L.ctx)
    out = {}
    for cid in problem.classes:
        need = sum(u.needed * u.length for u in problem.units if u.class_id == cid)
        out[problem.labels[cid]] = (need, len(problem.allowed[cid]), str(cid))
    return out


async def make_classes_fit(api: Api, over: dict[str, tuple[int, int, str]], subjects: list[dict]) -> dict[str, tuple[str, int, int]]:
    """The HOD's job: for each over-full class lower the weekly lectures of a few ordinary subjects (credits stay) until it fits with a little room."""
    changed: dict[str, tuple[str, int, int]] = {}
    for label, (need, free, _cid) in over.items():
        slack = 2
        cands = [s for s in subjects if s["class_label"] == label and not s["is_lab"] and not s["special_kind"] and not s["is_elective"] and s["weekly_effective"] > 1]
        weekly = {s["id"]: s["weekly_effective"] for s in cands}
        total = need
        while total > free - slack:
            pick = max((s for s in cands if weekly[s["id"]] > 1), key=lambda s: weekly[s["id"]], default=None)
            if pick is None:
                break
            weekly[pick["id"]] -= 1
            total -= 1
        for s in cands:
            if weekly[s["id"]] != s["weekly_effective"]:
                await api.patch(f"/api/subjects/{s['id']}", {"weekly_lectures": weekly[s["id"]]})
                changed[s["id"]] = (s["name"], s["credits"], weekly[s["id"]])
        log(f"   {label}: needed {need}, free {free} -> now needs {total}")
    return changed


async def scenarios(api: Api, client: httpx.AsyncClient) -> None:
    tree = await api.get("/api/structure/tree")
    sid = next(s for s in tree["schools"] if s["name"] == "School of IT")["id"]
    q = f"?school_id={sid}"
    today = vw.today_ist()
    monday = vw.monday_of(today)

    # ------------------------------------------------------------------------------------------------ 1) weekly lectures
    log("\n== 1) weekly-lectures override ==")
    nums = await class_numbers(sid)
    over = {k: v for k, v in nums.items() if v[0] > v[1]}
    check(set(over) == {"B.Sc. IT - Sem 3", "B.Sc. IT - Sem 5", "B.Sc. IT - Sem 7"}, "before: exactly the 3 over-full B.Sc. classes cannot fit their credits", str({k: f"{v[0]}/{v[1]}" for k, v in over.items()}))
    subjects = (await api.get(f"/api/subjects{q}"))["subjects"]
    check(all(s["weekly_lectures"] is None and s["weekly_effective"] == s["credits"] for s in subjects), "every seeded subject starts with no override (weekly lectures = credits)")
    changed = await make_classes_fit(api, over, subjects)
    after = {s["id"]: s for s in (await api.get(f"/api/subjects{q}"))["subjects"]}
    check(bool(changed) and all(after[i]["weekly_lectures"] == w and after[i]["credits"] == c for i, (_n, c, w) in changed.items()),
          f"{len(changed)} weekly-lectures values set by the HOD API; the credits did not change")
    nums2 = await class_numbers(sid)
    check(all(v[0] <= v[1] for v in nums2.values()), "after: every class now fits its day-wise week (precheck finds no capacity problem)", str({k: f"{v[0]}/{v[1]}" for k, v in nums2.items() if k in over}))
    sub_ids = {k: v for k, v in changed.items()}
    check(not any(s["is_lab"] or s["special_kind"] for s in after.values() if s["id"] in sub_ids), "labs and PE / SOL were never touched (a lab stays one double block)")

    t0 = monotonic()
    res = await api.post("/api/generation/runs", {"school_id": sid, "name": "P4 SELFTEST", "time_limit_s": 120})
    check(not [f for f in res["precheck"] if f["kind"] == "class_capacity"], "generation precheck: no 'class needs more periods than it has' finding any more")
    run = await wait_run(api, res["id"])
    check(run["status"] == "done", f"generation finished in {monotonic() - t0:.0f}s")
    best = run["solutions"][0]
    classes = best["breakdown"]["classes"]
    check(all(v["percent"] == 100 and v["missing"] == 0 for v in classes.values()), "the best solution places EVERY lecture of ALL 8 classes (100 %) - the B.Sc. 3 / 5 / 7 finding is resolved",
          str({v["label"]: v["percent"] for v in classes.values() if v["percent"] < 100}))
    check(not best["breakdown"]["unplaced"], "nothing is left unplaced")
    await api.post(f"/api/generation/solutions/{best['id']}/use")
    rep = await api.get(f"/api/conflicts{q}")
    check(rep["counts"]["clashes"] == 0 and rep["counts"]["missing"] == 0, "conflict report: 0 clashes, 0 missing lectures")
    listing = await api.get(f"/api/drafts?school_id={sid}")
    primaries = {c["label"]: next(d for d in c["drafts"] if d["is_primary"]) for c in listing["classes"]}
    sem_of = {c["label"]: c["semester_id"] for c in listing["classes"]}

    # ------------------------------------------------------------------------------------------------ 2) publish
    log("\n== 2) publish ==")
    teachers_all = {r["full_name"]: str(r["id"]) for r in await db.fetch("select id, full_name from public.profiles where school_id = $1 and role in ('teacher','hod')", __import__("uuid").UUID(sid))}
    st = await api.get(f"/api/publish/status{q}")
    check(st["summary"]["not_published"] == 8 and st["summary"]["can_publish"] == 8, "status: all 8 primary drafts are ready, none published yet")
    async with Acting("Mr. Sushant Singh") as _u:
        pre = await client.get("/api/views/my/week")
        pre_json = pre.json()
        check(pre.status_code == 200 and pre_json["published"] is False and pre_json["cells"] == [], "BEFORE publishing a teacher sees nothing (the drafts are invisible)")
        d403 = await client.get("/api/drafts")
        check(d403.status_code == 403, "a teacher cannot even list drafts (403)")
        d403b = await client.get(f"/api/drafts/{primaries['M.Sc. AI - Sem 1']['id']}")
        check(d403b.status_code == 403, "a teacher cannot open a draft by its id either (403)")

    one = await api.post("/api/publish", {"scope": "class", "semester_id": sem_of["M.Sc. AI - Sem 1"], "school_id": sid})
    check(one["published"][0]["version"] == 1 and len(one["published"]) == 1, "publishing ONE class works (version 1)")
    check(one["notified"] > 0, f"…and notified {one['notified']} teachers of that class only")
    first_ids = {r["user_id"] for r in await db.fetch("select user_id from public.notifications where kind = 'timetable_published'")}
    allp = await api.post("/api/publish", {"scope": "all", "school_id": sid})
    check(len(allp["published"]) == 7 and {p["class_label"] for p in allp["published"]} == set(sem_of) - {"M.Sc. AI - Sem 1"}, "publish ALL: the 7 remaining classes go live; the already-live class is not re-published")
    check(all(p["version"] == 1 for p in allp["published"]), "every class is at version 1")
    notes = await db.fetch("select user_id, kind, dedupe_key from public.notifications")
    first_notes = [n for n in notes if n["kind"] == "timetable_published"]
    by_user_event = {(n["user_id"], n["dedupe_key"]) for n in first_notes}
    check(len(by_user_event) == len(first_notes), "never two notifications for the same person from the same publish click")
    with_lectures = {r["teacher_id"] for r in await db.fetch("select distinct teacher_id from public.published_entries where teacher_id is not null")} | {
        r["co_teacher_id"] for r in await db.fetch("select distinct co_teacher_id from public.published_entries where co_teacher_id is not null")}
    notified_users = {n["user_id"] for n in first_notes}
    check(notified_users == with_lectures, f"exactly the {len(with_lectures)} people who teach something were notified - nobody else", f"notified {len(notified_users)}")
    audits = await db.fetch("select details from public.audit_log where entity = 'publication' and at >= (now() - interval '10 minutes')")
    check(len(audits) == 8, "the audit log records one 'publish' line per class with version and counts")
    again = await client.post("/api/publish", json={"scope": "all", "school_id": sid})
    check(again.status_code == 409 and again.json()["code"] == "NOTHING_TO_PUBLISH", "publishing again with no change is refused (no new version, no spam)")
    status2 = await api.get(f"/api/publish/status{q}")
    check(status2["summary"]["in_sync"] == 8 and status2["summary"]["unpublished"] == 0, "status: all 8 up to date")

    # -- change ONE lecture in ONE class's primary draft -> "unpublished changes" -> re-publish tells only its teacher(s)
    target_label, draft_id, mover, moved = await find_move(api, client, primaries, prefer="M.Sc. IT - Sem 3")
    check(moved is not None, f"a lecture of {target_label} was moved in the primary draft (to day {moved[0]} period {moved[1]})")
    mover_teachers = {e["teacher_id"] for e in mover["entries"] if e["teacher_id"]} | {e["co_teacher_id"] for e in mover["entries"] if e["co_teacher_id"]}
    status3 = await api.get(f"/api/publish/status{q}")
    ch = {c["class_label"]: c for c in status3["classes"]}
    check([k for k, c in ch.items() if c["state"] == "changed"] == [target_label], "ONLY that class shows 'unpublished changes'; the 7 others are up to date")
    check(ch[target_label]["changes"]["added_count"] == 1 and ch[target_label]["changes"]["removed_count"] == 1, "…with the exact difference: 1 lecture added at its new time, 1 removed from its old time")
    # a teacher still sees the OLD timetable
    owner_name = next(k for k, v in teachers_all.items() if v in {str(t) for t in mover_teachers})
    async with Acting(owner_name):
        old_week = (await client.get("/api/views/my/week")).json()
    old_cell = [c for c in old_week["cells"] if c["items"][0]["title"] in {e["subject_name"] for e in mover["entries"]} and (c["day"], c["start_period"]) == (mover["day"], mover["start_period"])]
    check(len(old_cell) == 1, f"before re-publishing, {owner_name} still sees the lecture at its OLD time (a draft change is invisible)")
    before_notes = {str(r["user_id"]): r["n"] for r in await db.fetch("select user_id, count(*) n from public.notifications group by user_id")}
    rp = await api.post("/api/publish", {"scope": "all", "school_id": sid})
    check([p["class_label"] for p in rp["published"]] == [target_label] and rp["published"][0]["version"] == 2, "re-publish: only the changed class goes to version 2")
    after_notes = {str(r["user_id"]): r["n"] for r in await db.fetch("select user_id, count(*) n from public.notifications group by user_id")}
    got_new = {u for u in after_notes if after_notes[u] > before_notes.get(u, 0)}
    check(got_new == {str(t) for t in mover_teachers}, f"ONLY the teacher(s) whose own lecture changed were notified: {sorted(teachers_all)[:0] or [k for k, v in teachers_all.items() if v in got_new]}")
    check(rp["notified"] == len(mover_teachers), "and exactly one notification each")
    changed_note = await db.fetchrow("select title, body from public.notifications where kind = 'timetable_changed' order by created_at desc limit 1")
    check(changed_note and "moved" in changed_note["body"], "the message says what happened", (changed_note or {}).get("body", "")[:110])
    async with Acting(owner_name):
        new_week = (await client.get("/api/views/my/week")).json()
    new_cell = [c for c in new_week["cells"] if c["items"][0]["title"] in {e["subject_name"] for e in mover["entries"]} and (c["day"], c["start_period"]) == moved]
    check(len(new_cell) == 1, f"after re-publishing {owner_name} sees it at the NEW time")

    # ------------------------------------------------------------------------------------------------ 3) what teachers see
    log("\n== 3) teacher views ==")
    async with Acting("Mr. Sushant Singh"):
        w = (await client.get(f"/api/views/my/week?week_of={monday}")).json()
        pe = [c for c in w["cells"] if c["kind"] == "fixed" and c["fixed_kind"] == "PE" and c["combined"]]
        check(len(pe) == 1 and pe[0]["combined"] and len(pe[0]["classes"]) == 5, "the COMBINED PE appears ONCE in his week with all 5 classes", str([c["label"] for c in pe[0]["classes"]]) if pe else "")
        check(len(pe[0]["items"]) == 1, "…as one session (not 5 separate lectures)")
        wd = today.isoweekday()
        t = (await client.get("/api/views/my/today")).json()
        expect = [c for c in w["cells"] if c["day"] == wd]
        check(t["weekday"] == wd and t["date"] == today.isoformat() and len(t["cells"]) == len(expect), f"Home = today's own lectures ({t['weekday_name']}): {len(t['cells'])} lectures")
        sun = (await client.get(f"/api/views/my/today?on={monday + timedelta(days=6)}")).json()
        check(sun["weekday_name"] == "Sunday" and sun["is_working_day"] is False and sun["cells"] == [], "a Sunday is correctly 'not a working day'")
    async with Acting("Dr. Papri Das") as _a:
        wa = (await client.get(f"/api/views/my/week?week_of={monday}")).json()
    async with Acting("Maitri Hingu"):
        wb = (await client.get(f"/api/views/my/week?week_of={monday}")).json()
    qa = [c for c in wa["cells"] if c["items"][0]["title"] == "Introduction to Quantum Computing"]
    qb = [c for c in wb["cells"] if c["items"][0]["title"] == "Introduction to Quantum Computing"]
    check(len(qa) >= 1 and {(c["day"], c["start_period"]) for c in qa} == {(c["day"], c["start_period"]) for c in qb}, "a CO-TAUGHT subject (Papri Das + Maitri Hingu) appears in BOTH teachers' weeks at the same times")
    async with Acting("External Faculty (AI&ML)"):
        wx = (await client.get(f"/api/views/my/week?week_of={monday}")).json()
    par = [c for c in wx["cells"] if c["parallel_with"] if "parallel_with" in c]
    check(len(par) >= 1 and all(len(c["items"]) == 1 for c in par), "an elective-track teacher sees only their own option, and what runs in parallel with it", str(par[0]["parallel_with"]) if par else "")

    # holiday: shown on its day for the teacher, lecture marked cancelled
    pick = next(c for c in wx["cells"] if c["kind"] == "lecture")
    hol_date = monday + timedelta(days=7 + pick["day"] - 1)
    hol = await api.post("/api/calendar", {"scope": "school", "school_id": sid, "kind": "holiday", "title": "ZZ SELFTEST Holiday", "date_from": hol_date.isoformat()})
    async with Acting("External Faculty (AI&ML)"):
        th = (await client.get(f"/api/views/my/today?on={hol_date}")).json()
    check(th["banners"] and th["banners"][0]["title"] == "ZZ SELFTEST Holiday" and any(c["status"] == "holiday" for c in th["cells"]), "a holiday is shown on that day and his lectures are marked 'not happening'")
    check(th["stats"]["cancelled"] >= 1, "…and counted as cancelled, not as teaching hours")
    await api.call("DELETE", f"/api/calendar/{hol['id']}")

    # a draft the HOD is still working on never reaches a teacher
    extra_draft = await api.post("/api/drafts", {"semester_id": sem_of["M.Sc. IT - Sem 1"], "name": "SELFTEST secret draft", "copy_from": primaries["M.Sc. IT - Sem 1"]["id"]})
    secret = await api.get(f"/api/drafts/{extra_draft['id']}")
    victim = secret["placements"][0]
    await client.delete(f"/api/drafts/{extra_draft['id']}/placements/{victim['key']}")
    async with Acting("Dr. Juhi Khengar"):
        wj = (await client.get(f"/api/views/my/week?week_of={monday}")).json()
        mj = await client.get("/api/views/master")
    check(wj["published"] and wj["stats"]["lectures"] > 0, "teachers keep seeing the published timetable while other drafts exist")
    check(mj.status_code == 403, "a teacher cannot call the school-wide master view (403)")
    master = await api.get(f"/api/views/master?school_id={sid}&scope=week&week_of={monday}")
    n_master = master["stats"]["lectures"]
    published_rows = await db.fetchval("select count(distinct coalesce(group_key, id)) from public.published_entries")
    check(n_master == published_rows, "the master view equals the published rows exactly (drafts are never mixed in)", f"{n_master} lectures")
    await client.delete(f"/api/drafts/{extra_draft['id']}")

    # HOD / admin views and filters
    opts = await api.get(f"/api/views/options{q}")
    check(len(opts["classes"]) == 8 and all(c["published"] for c in opts["classes"]), "class-wise view: all 8 classes are selectable and published")
    cv = await api.get(f"/api/views/class/{sem_of['B.Sc. IT - Sem 5']}?week_of={monday}")
    check(cv["stats"]["lectures"] > 0 and all(len(c["classes"]) == 1 for c in cv["cells"]), "class-wise view: B.Sc. IT - Sem 5 week")
    tv = await api.get(f"/api/views/teacher/{teachers_all['Dr. Chirag Mehta']}?week_of={monday}")
    check(tv["stats"]["lectures"] > 0, "teacher-wise view: Dr. Chirag Mehta's week", str(tv["stats"]))
    mf = await api.get(f"/api/views/master?school_id={sid}&scope=week&week_of={monday}&semester_id={sem_of['M.Sc. AI - Sem 1']}&teacher_id={teachers_all['Anjana Prajapati']}&day=1")
    check(all(c["day"] == 1 for c in mf["cells"]) and all(any(x["label"] == "M.Sc. AI - Sem 1" for x in c["classes"]) for c in mf["cells"]) and mf["classes"][0]["label"] == "M.Sc. AI - Sem 1", "master view filters (class + teacher + day) combine correctly")
    mt = await api.get(f"/api/views/master?school_id={sid}&scope=today")
    check(mt["date"] == today.isoformat() and all(c["day"] == today.isoweekday() for c in mt["cells"]), "master 'today' filter shows only today's lectures")

    # ------------------------------------------------------------------------------------------------ 4) exports
    log("\n== 4) PDF + Excel exports ==")
    from openpyxl import load_workbook
    from pypdf import PdfReader

    def pdf_text(raw: bytes) -> str:
        return "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(raw)).pages)

    r = await client.get(f"/api/exports/class/{sem_of['B.Sc. IT - Sem 3']}?format=pdf")
    t = pdf_text(r.content)
    check(r.status_code == 200 and r.content[:5] == b"%PDF-" and len(PdfReader(io.BytesIO(r.content)).pages) == 1, "class PDF: one clean A4 page", r.headers.get("content-disposition", ""))
    check(all(x in t for x in ("AURO University", "School of IT", "B.Sc. IT - Sem 3", "Odd 2026", "LUNCH", "13:10-13:40", "break")), "…with college, school, class, term, lunch and break rows")
    check("Combined with" in t and ("PARALLEL" in t or "Parallel" in t), "…combined sessions and parallel electives are labelled")
    r = await client.get(f"/api/exports/class/{sem_of['B.Sc. IT - Sem 3']}?format=xlsx")
    wb = load_workbook(io.BytesIO(r.content))
    check(r.status_code == 200 and wb["Timetable"].page_setup.orientation == "landscape" and "Subjects" in wb.sheetnames and wb["Timetable"].merged_cells.ranges, "class Excel: landscape, print-ready, merged blocks, subject sheet")
    tid = teachers_all["Mr. Sushant Singh"]
    r = await client.get(f"/api/exports/teacher/{tid}?format=pdf")
    t = pdf_text(r.content)
    check(r.status_code == 200 and "Teacher Timetable - Mr. Sushant Singh" in t and "Combined session" in t and "Physical Education" in t, "teacher PDF: his week with the combined PE session")
    r = await client.get(f"/api/exports/teacher/{tid}?format=xlsx")
    check(r.status_code == 200 and "Lectures" in load_workbook(io.BytesIO(r.content)).sheetnames, "teacher Excel")
    r = await client.get(f"/api/exports/master?school_id={sid}&format=pdf")
    reader = PdfReader(io.BytesIO(r.content))
    t = pdf_text(r.content)
    check(r.status_code == 200 and len(reader.pages) == 6 and all(c in t for c in ("B.Sc. IT - Sem 1", "M.Sc. IT - Sem 3", "School Master Timetable - Saturday")), "school master PDF: one page per weekday (Mon-Sat), every class", f"{len(reader.pages)} pages")
    r = await client.get(f"/api/exports/master?school_id={sid}&format=xlsx")
    wb = load_workbook(io.BytesIO(r.content))
    check(r.status_code == 200 and wb.sheetnames[:2] == ["Monday", "Tuesday"] and "All lectures" in wb.sheetnames, "school master Excel: a sheet per weekday + a flat 'All lectures' sheet", str(wb.sheetnames))
    r = await client.get(f"/api/exports/master?school_id={sid}&format=pdf&semester_id={sem_of['M.Sc. AI - Sem 1']}&day=1")
    check(r.status_code == 200 and len(PdfReader(io.BytesIO(r.content)).pages) == 1, "filtered master PDF (one class, one day) is a single page")

    # ------------------------------------------------------------------------------------------------ 5) notifications
    log("\n== 5) notifications: bell + live feed + Web Push ==")
    probe_name = "Dr. Dhaval Thaker"
    probe_id = teachers_all[probe_name]
    listing = None
    async with Acting(probe_name):
        listing = (await client.get("/api/notifications")).json()
        check(listing["unread"] >= 1 and listing["notifications"][0]["link"] == "/my/timetable", "the bell lists the person's own notifications with unread count and a link", listing["notifications"][0]["title"])
        nid = listing["notifications"][0]["id"]
        r1 = (await client.post("/api/notifications/read", json={"ids": [nid]})).json()
        check(r1["unread"] == listing["unread"] - 1, "mark one read lowers the unread count")
        r2 = (await client.post("/api/notifications/read", json={"all": True})).json()
        check(r2["unread"] == 0, "mark all read -> 0 unread")
    # live feed: subscribe to the in-process hub like an open bell, then cause a change
    q_hub = notify.hub.subscribe(probe_id)
    pe_status = (await api.get(f"/api/publish/status{q}"))["summary"]
    check(pe_status["unpublished"] == 0, "(everything is published; now one more change will be made)")
    # a local "browser" that receives Web Push
    ec_key, auth, p256dh = new_browser_key()
    server = start_push_server()
    notify.ALLOW_LOCAL_ENDPOINTS = True
    await db.execute(
        "insert into public.push_subscriptions (user_id, endpoint, p256dh, auth, user_agent) values ($1, $2, $3, $4, 'selftest-browser')",
        __import__("uuid").UUID(probe_id), f"http://127.0.0.1:{server.server_port}/push", p256dh, auth_b64(auth),
    )
    # move a lecture of the probe teacher and publish
    my_label, my_draft, my_mover, moved2 = await find_move(api, client, primaries, teacher_id=probe_id)
    check(moved2 is not None, f"moved one of {probe_name}'s lectures in {my_label}")
    rp2 = await api.post("/api/publish", {"scope": "class", "semester_id": sem_of[my_label], "school_id": sid})
    check(rp2["notified"] >= 1, "re-publish sent a notification to the affected teacher(s)")
    try:
        live_item = await asyncio.wait_for(q_hub.get(), timeout=5)
        check(live_item["kind"] == "timetable_changed" and live_item["is_read"] is False, "LIVE: the open bell received the notification at once (no reload)", live_item["title"])
    except asyncio.TimeoutError:
        check(False, "LIVE: the open bell received the notification at once")
    notify.hub.unsubscribe(probe_id, q_hub)
    await notify.drain()
    check(len(PUSH_ARRIVED) == 1, "WEB PUSH: the 'browser' received exactly ONE push request", str(len(PUSH_ARRIVED)))
    if PUSH_ARRIVED:
        import http_ece

        got = PUSH_ARRIVED[0]
        plain = json.loads(http_ece.decrypt(got["body"], private_key=ec_key, auth_secret=auth))
        hdr = {k.lower(): v for k, v in got["headers"].items()}
        check(hdr.get("content-encoding") == "aes128gcm" and hdr.get("authorization", "").lower().startswith("vapid "), "…encrypted (aes128gcm) and signed with our VAPID key")
        check(plain["title"] == "Your timetable has changed" and plain["url"] == "/my/timetable" and "moved" in plain["body"], "…and only that browser can read it: 'Your timetable has changed' - " + plain["body"][:70])
    sub_count = await db.fetchval("select count(*) from public.push_subscriptions where user_id = $1", __import__("uuid").UUID(probe_id))
    check(sub_count == 1, "a working subscription is kept")
    PUSH_ARRIVED.clear()
    server.respond = 410
    my_draft2 = my_draft
    r = await client.patch(f"/api/drafts/{my_draft2}/placements/{my_mover['key']}", json={"day": my_mover["day"], "start_period": my_mover["start_period"]})
    if r.status_code == 200:
        await api.post("/api/publish", {"scope": "class", "semester_id": sem_of[my_label], "school_id": sid})
        await notify.drain()
    check(len(PUSH_ARRIVED) == 1 and await db.fetchval("select count(*) from public.push_subscriptions where user_id = $1", __import__("uuid").UUID(probe_id)) == 0,
          "a browser that answers '410 gone' is removed automatically (no endless retries)")
    server.shutdown()
    notify.ALLOW_LOCAL_ENDPOINTS = False
    unread_for_nobody = await db.fetchval("select count(*) from public.notifications n join public.profiles p on p.id = n.user_id where not p.is_active")
    check(unread_for_nobody == 0, "nobody deactivated received anything")

    # ------------------------------------------------------------------------------------------------ 6) hours + compensation
    log("\n== 6) semester hours + holiday compensation ==")
    base = await api.get(f"/api/hours/subjects?school_id={sid}")
    s0 = base["summary"]
    check(s0["subjects"] == 93, "93 subjects are tracked from the published timetable + semester dates (1 Jul - 30 Nov)", f"{s0['short_subjects']} short")
    mai = next(c for c in base["classes"] if c["class_label"] == "M.Sc. AI - Sem 1")
    check(mai["dates"] == {"start": "2026-07-01", "end": "2026-11-30"} and mai["published"], "M.Sc. AI - Sem 1: semester dates + published status are read correctly")
    brk = await api.post("/api/calendar", {"scope": "school", "school_id": sid, "kind": "holiday", "title": "ZZ SELFTEST Break", "date_from": "2026-07-01", "date_to": "2026-08-30", "semester_ids": [sem_of["M.Sc. AI - Sem 1"]]})
    h1 = await api.get(f"/api/hours/subjects?school_id={sid}&semester_id={sem_of['M.Sc. AI - Sem 1']}")
    mai1 = h1["classes"][0]
    check(mai1["short_subjects"] > 0 and h1["summary"]["short_hours"] > 0, f"the new holiday immediately shows RED shortfalls: {mai1['short_subjects']} subjects short, {mai1['short_hours']} hours")
    tgt = max(mai1["subjects"], key=lambda s: s["shortfall"])
    check(tgt["shortfall"] > 0 and tgt["lost"] > 0, f"'{tgt['name']}' is short by {tgt['shortfall']} hours (lost {tgt['lost']} h to the holiday)")
    async with Acting(tgt["teachers"][0]["name"]):
        mine = (await client.get("/api/hours/my")).json()
    check(any(s["subject_id"] == tgt["subject_id"] and s["shortfall"] > 0 for s in mine["subjects"]), "the teacher sees the same red 'short by' on My Subjects")
    dash = await api.get(f"/api/hours/summary?school_id={sid}")
    check(dash["schools"][0]["summary"]["short_subjects"] >= mai1["short_subjects"], "the dashboard summary counts it too")
    sug = await api.post("/api/hours/suggest", {"subject_id": tgt["subject_id"]})
    check(len(sug["suggestions"]) > 0, f"one click suggests {len(sug['suggestions'])} clash-free extra slot(s)", sug["message"][:140])
    # an independent clash check against the published timetable
    rows = await db.fetch("select semester_id, day, start_period, length, teacher_id, co_teacher_id, group_key, id from public.published_entries")
    sem_id = __import__("uuid").UUID(sem_of["M.Sc. AI - Sem 1"])
    sub_row = await db.fetchrow("select id, semester_id from public.subjects where id = $1", __import__("uuid").UUID(tgt["subject_id"]))
    teachers_of_subject = {__import__("uuid").UUID(t["id"]) for t in tgt["teachers"]}
    hol_from, hol_to = date(2026, 7, 1), date(2026, 8, 30)
    bad_slots = []
    for s in sug["suggestions"]:
        d = date.fromisoformat(s["date"])
        wd = d.isoweekday()
        length = s["length"]
        clash_class = any(r["semester_id"] == sem_id and r["day"] == wd and r["start_period"] < s["start_period"] + length and s["start_period"] < r["start_period"] + r["length"] for r in rows)
        clash_teacher = any((teachers_of_subject & {r["teacher_id"], r["co_teacher_id"]}) and r["day"] == wd and r["start_period"] < s["start_period"] + length and s["start_period"] < r["start_period"] + r["length"] for r in rows)
        wrong_day = d <= today or wd > 6 or not date(2026, 7, 1) <= d <= date(2026, 11, 30) or hol_from <= d <= hol_to
        if clash_class or clash_teacher or wrong_day:
            bad_slots.append((s["date"], s["start_period"], clash_class, clash_teacher, wrong_day))
    check(not bad_slots, "every suggested slot is verified clash-free against the whole published school (class, teacher, co-teacher, working day, no holiday, inside the semester)", str(bad_slots[:2]))
    pick2 = sug["suggestions"][:2]
    slots = [{"date": s["date"], "start_period": s["start_period"]} for s in pick2]
    appr = await api.post("/api/hours/extra", {"subject_id": tgt["subject_id"], "slots": slots, "reason": "SELFTEST compensation"})
    check(appr["lectures"] == len(slots) and appr["notified"] >= 1, "the HOD approves -> extra one-off lectures are added and the teacher is notified", appr["message"])
    h2 = await api.get(f"/api/hours/subjects?school_id={sid}&semester_id={sem_of['M.Sc. AI - Sem 1']}")
    t2 = next(s for s in h2["classes"][0]["subjects"] if s["subject_id"] == tgt["subject_id"])
    check(t2["extra"] == tgt["extra"] + len(slots) and t2["shortfall"] == tgt["shortfall"] - len(slots) * (2 if t2["is_lab"] else 1), "the shortfall shrank by exactly the approved hours")
    first_slot = date.fromisoformat(pick2[0]["date"])
    async with Acting(tgt["teachers"][0]["name"]):
        wk = (await client.get(f"/api/views/my/week?week_of={vw.monday_of(first_slot)}")).json()
        nn = (await client.get("/api/notifications")).json()
    extras_cells = [c for c in wk["cells"] if c["kind"] == "extra" and c["date"] == pick2[0]["date"]]
    check(len(extras_cells) == 1 and extras_cells[0]["start_period"] == pick2[0]["start_period"], "the extra lecture shows in the teacher's own week on that exact date")
    check(any(n["title"].lower().startswith(("extra lecture", "2 extra")) for n in nn["notifications"]), "…and a notification 'extra lecture added' reached the teacher")
    # the final state, with the extras in place, is still free of double-booking
    ex = await db.fetch("select semester_id, lecture_date, start_period, length, teacher_id, co_teacher_id from public.extra_lectures")
    doubles = 0
    for i, a in enumerate(ex):
        for b in ex[i + 1:]:
            if a["lecture_date"] == b["lecture_date"] and a["start_period"] < b["start_period"] + b["length"] and b["start_period"] < a["start_period"] + a["length"] and (
                a["semester_id"] == b["semester_id"] or ({a["teacher_id"], a["co_teacher_id"]} & {b["teacher_id"], b["co_teacher_id"]} - {None})):
                doubles += 1
    check(doubles == 0, "the approved extra lectures do not collide with each other")
    dup = await client.post("/api/hours/extra", json={"subject_id": tgt["subject_id"], "slots": slots[:1]})
    check(dup.status_code == 422, "approving the same slot twice is refused (it is taken now)")
    await api.call("DELETE", f"/api/calendar/{brk['id']}")


# ====================================================================================================================
def new_browser_key():
    from cryptography.hazmat.primitives import serialization
    from cryptography.hazmat.primitives.asymmetric import ec

    key = ec.generate_private_key(ec.SECP256R1())
    auth = os.urandom(16)
    b64 = lambda raw: base64.urlsafe_b64encode(raw).decode().rstrip("=")  # noqa: E731
    p256dh = b64(key.public_key().public_bytes(serialization.Encoding.X962, serialization.PublicFormat.UncompressedPoint))
    return key, auth, p256dh


def auth_b64(raw: bytes) -> str:
    return base64.urlsafe_b64encode(raw).decode().rstrip("=")


def start_push_server() -> HTTPServer:
    class Handler(BaseHTTPRequestHandler):
        def do_POST(self):  # noqa: N802
            n = int(self.headers.get("Content-Length", 0))
            PUSH_ARRIVED.append({"body": self.rfile.read(n), "headers": dict(self.headers)})
            self.send_response(self.server.respond)  # type: ignore[attr-defined]
            self.end_headers()

        def log_message(self, *a):
            pass

    server = HTTPServer(("127.0.0.1", 0), Handler)
    server.respond = 201  # type: ignore[attr-defined]
    threading.Thread(target=server.serve_forever, daemon=True).start()
    return server


async def find_move(api: Api, client: httpx.AsyncClient, primaries: dict, prefer: str | None = None, teacher_id: str | None = None):
    """Moves one single-period, unpinned lecture to a cell the clash engine allows. Returns (class label, draft id, placement, (day, start))."""
    order = sorted(primaries, key=lambda k: (k != prefer, k))
    for label in order:
        draft = primaries[label]["id"]
        det = await api.get(f"/api/drafts/{draft}")
        for p in det["placements"]:
            if p["parallel"] or p["is_locked"] or p["length"] != 1:
                continue
            if teacher_id and not any(e["teacher_id"] == teacher_id for e in p["entries"]):
                continue
            opts = await api.post(f"/api/drafts/{draft}/options", {"placement_key": p["key"]})
            for key, cell in opts["cells"].items():
                day, start = (int(x) for x in key.split(":"))
                if cell["state"] in ("ok", "warn") and (day, start) != (p["day"], p["start_period"]):
                    r = await client.patch(f"/api/drafts/{draft}/placements/{p['key']}", json={"day": day, "start_period": start})
                    if r.status_code == 200:
                        return label, draft, p, (day, start)
    raise SeedError("no movable lecture found")


async def cleanup(api: Api, started, original: dict[str, int | None]) -> None:
    log("\n== cleanup: every test artifact goes away ==")
    import uuid

    tree = await api.get("/api/structure/tree")
    sid = next(s for s in tree["schools"] if s["name"] == "School of IT")["id"]
    school = uuid.UUID(sid)
    await db.execute("delete from public.extra_lectures where school_id = $1", school)
    await db.execute("delete from public.calendar_entries where title like 'ZZ SELFTEST%'")
    await db.execute("delete from public.notifications where created_at >= $1", started)
    await db.execute("delete from public.push_subscriptions where user_agent = 'selftest-browser'")
    await db.execute("delete from public.publish_events where school_id = $1", school)
    await db.execute("delete from public.publications where school_id = $1", school)           # published_entries go with them
    listing = await api.get(f"/api/drafts?school_id={sid}")
    removed = 0
    for c in listing["classes"]:
        for d in c["drafts"]:
            await api.call("DELETE", f"/api/drafts/{d['id']}")
            removed += 1
    runs = await api.get(f"/api/generation/runs?school_id={sid}")
    for r in runs["runs"]:
        if r["status"] == "running":
            await api.post(f"/api/generation/runs/{r['id']}/cancel")
            await wait_run(api, r["id"])
        await api.call("DELETE", f"/api/generation/runs/{r['id']}")
    await db.execute("delete from public.generation_settings where school_id = $1", school)
    restored = 0
    for sid_, weekly in original.items():
        n = await db.execute("update public.subjects set weekly_lectures = $2 where id = $1 and weekly_lectures is distinct from $2", uuid.UUID(sid_), weekly)
        restored += int(n.split()[-1])
    await db.execute("delete from public.audit_log where at >= $1", started)
    left = await db.fetchrow(
        """
        select (select count(*) from public.timetable_drafts) d, (select count(*) from public.generation_runs) g,
               (select count(*) from public.publications) p, (select count(*) from public.published_entries) pe,
               (select count(*) from public.publish_events) ev, (select count(*) from public.notifications) n,
               (select count(*) from public.push_subscriptions) ps, (select count(*) from public.extra_lectures) x,
               (select count(*) from public.subjects where weekly_lectures is not null) w,
               (select count(*) from public.calendar_entries where title like 'ZZ%') h
        """
    )
    check(all(v == 0 for v in dict(left).values()) and removed > 0,
          f"{removed} drafts, every run, publication, event, notification, push subscription, extra lecture, test holiday and {restored} weekly-lectures override removed", str(dict(left)))


if __name__ == "__main__":
    sys.exit(asyncio.run(main()))
