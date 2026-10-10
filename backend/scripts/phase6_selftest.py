"""Phase 6 self-test on the REAL seeded School of IT data (in-process API as the real admin; teachers / the HOD are impersonated).

  python -m scripts.phase6_selftest

  1. a baseline generation, so the "not available" cells chosen below are cells the teachers would REALLY have been given
  2. three real teachers (one of them the HOD) fill their preferences through their own screens: limit, visiting days, special choices
  3. the preference window: open (who is told), live filled / not filled, reminder (once, only the unfilled), auto-close, reopen, extend, close
  4. generation WITH the preferences: no "not available" slot and no daily maximum is ever crossed; the star scoring is in the breakdown
  5. the manual grid refuses a drop onto a "not available" cell and past a daily maximum, with the reason
  6. publish, then every report is checked against independent arithmetic on the published rows (and the Excel files)
  7. an HOD is locked to his own school in every report
  8. the audit screen shows the real history, filtered and paginated
  9. cleanup: every preference / window / draft / publication / notification is removed, the preferences are restored, the seed is verified again
"""
from __future__ import annotations

import os

os.environ.setdefault("READ_CACHE_SECONDS", "0")   # these scripts also write straight to the database: always read fresh

import asyncio
import io
import sys
import uuid
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta, timezone

import httpx
from openpyxl import load_workbook

from app import db
from app.config import get_settings
from app.main import app
from app.security import deps
from app.services import notify, prefflow
from app.services import views as vw

from .phase3_selftest import wait_run
from .phase4_selftest import Acting, class_numbers, make_classes_fit
from .seed_college import Api, SeedError, Who, _fake_authenticate, load_user, log, verify

OK = BAD = 0
HOD_NAME = "Dr. Papri Das"


def check(cond: bool, msg: str, extra: str = "") -> None:
    global OK, BAD
    OK += bool(cond)
    BAD += not cond
    log(("  PASS  " if cond else "  FAIL  ") + msg + (f"   {extra}" if extra else ""))


def nxt(wd: int) -> date:
    base = vw.today_ist()
    return base + timedelta(days=(wd - base.isoweekday()) % 7 or 7)


async def main() -> int:
    settings = get_settings()
    async with app.router.lifespan_context(app):
        started = await db.fetchval("select now()")
        Who.user = await load_user(email=settings.seed_admin_email, role="admin")
        deps.authenticate = _fake_authenticate
        async with httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://localhost",
                                     headers={"Origin": "http://localhost:5173"}, timeout=300) as client:
            api = Api(client)
            original = {str(r["id"]): r["weekly_lectures"] for r in await db.fetch("select id, weekly_lectures from public.subjects")}
            before = await snapshot()
            try:
                await scenarios(api, client, started)
            except (SeedError, AssertionError) as exc:
                log(f"\nSTOPPED: {exc}")
                return 1
            finally:
                await notify.drain()
                await cleanup(api, started, original, before)
            log("\n== the seeded data is untouched ==")
            bad = await verify(api)
            check(bad == 0, "seed verifier: ALL PASS")
    log(f"\nPHASE 6 SELF-TEST: {OK} checks passed, {BAD} failed")
    return 1 if BAD else 0


async def snapshot() -> dict:
    """What the preference tables looked like before the test (restored at the end)."""
    return {
        "prefs": [dict(r) for r in await db.fetch("select * from public.teacher_preferences")],
        "slots": [dict(r) for r in await db.fetch("select * from public.teacher_pref_slots")],
        "settings": [dict(r) for r in await db.fetch("select * from public.preference_settings")],
        "norms": [dict(r) for r in await db.fetch("select * from public.workload_norms")],
        "windows": await db.fetchval("select count(*) from public.preference_windows"),
    }


# ====================================================================================================
async def teacher_cells(solution_id: str, school: uuid.UUID) -> dict[uuid.UUID, list[tuple[int, int]]]:
    """Every (day, period) each teacher teaches in a generated solution (co-teachers and elective options included), plus locked sessions."""
    rows = await db.fetch(
        "select e.day, e.start_period, e.length, e.subject_id from public.draft_entries e join public.timetable_drafts d on d.id = e.draft_id where d.solution_id = $1",
        uuid.UUID(solution_id),
    )
    assign = {r["subject_id"]: (r["teacher_id"], r["co_teacher_id"]) for r in await db.fetch("select subject_id, teacher_id, co_teacher_id from public.assignments where school_id = $1", school)}
    out: dict[uuid.UUID, set[tuple[int, int]]] = defaultdict(set)
    for r in rows:
        for t in assign.get(r["subject_id"], ()):
            if t is not None:
                out[t].update((r["day"], r["start_period"] + i) for i in range(r["length"]))
    for f in await db.fetch("select day, start_period, length, teacher_id, co_teacher_id from public.fixed_slots where school_id = $1", school):
        for t in (f["teacher_id"], f["co_teacher_id"]):
            if t is not None:
                out[t].update((f["day"], f["start_period"] + i) for i in range(f["length"]))
    return {t: sorted(c) for t, c in out.items()}


async def notes(user_id: uuid.UUID, kind: str) -> int:
    return await db.fetchval("select count(*) from public.notifications where user_id = $1 and kind = $2", user_id, kind)


async def scenarios(api: Api, client: httpx.AsyncClient, started) -> None:
    tree = await api.get("/api/structure/tree")
    sid = next(s for s in tree["schools"] if s["name"] == "School of IT")["id"]
    school = uuid.UUID(sid)
    q = f"?school_id={sid}"
    hod_id = await db.fetchval("select id from public.profiles where full_name = $1", HOD_NAME)
    admin_id = Who.user.id

    async def as_(name: str, method: str, path: str, **kw) -> httpx.Response:
        async with Acting(name):
            return await client.request(method, path, **kw)

    async def run_generation(name: str, weights: dict | None = None) -> dict:
        body: dict = {"school_id": sid, "name": name, "time_limit_s": 120}
        if weights is not None:
            body["weights"] = weights
        res = await api.post("/api/generation/runs", body)
        run = await wait_run(api, res["id"])
        check(run["status"] == "done", f"'{name}' finished", f"{run['solutions'][0]['score']} points" if run.get("solutions") else str(run.get("message")))
        return run

    # ------------------------------------------------------------------------------------------------ 1) a baseline
    log("\n== 1) setup: weekly lectures that fit, then a baseline generation ==")
    nums = await class_numbers(sid)
    over = {k: v for k, v in nums.items() if v[0] > v[1]}
    subjects = (await api.get(f"/api/subjects{q}"))["subjects"]
    await make_classes_fit(api, over, subjects)
    base = await run_generation("P6 baseline")
    base_best = base["solutions"][0]
    base_cells = await teacher_cells(base_best["id"], school)

    load = await db.fetch(
        "select a.teacher_id, p.full_name, count(*) n from public.assignments a join public.profiles p on p.id = a.teacher_id join public.subjects s on s.id = a.subject_id "
        "where s.school_id = $1 and s.special_kind is null and not p.is_visiting and p.role = 'teacher' group by 1, 2 order by n desc, p.full_name", school,
    )
    t1_id, t2_id = load[0]["teacher_id"], load[1]["teacher_id"]
    t1, t2 = load[0]["full_name"], load[1]["full_name"]
    t3 = HOD_NAME
    log(f"   chosen: {t1} ({load[0]['n']} subjects), {t2} ({load[1]['n']}), and the HOD {t3} as a teacher")

    def spread(cells: list[tuple[int, int]], n: int) -> list[tuple[int, int]]:
        """n cells on different days if possible (the ones this teacher really got in the baseline)."""
        chosen: list[tuple[int, int]] = []
        for day in sorted({d for d, _p in cells}):
            chosen.append(next(c for c in cells if c[0] == day))
        return chosen[:n]

    def free_cells(taken: list[tuple[int, int]], n: int, days=(1, 2, 3)) -> list[tuple[int, int]]:
        return [(d, p) for d in days for p in (2, 3, 4) if (d, p) not in taken][:n]

    plan = {
        t1: {"id": t1_id, "no": spread(base_cells[t1_id], 4), "max": max(Counter(d for d, _p in base_cells[t1_id]).values()), "first": False, "last": True},
        t2: {"id": t2_id, "no": spread(base_cells[t2_id], 3), "max": None, "first": True, "last": False},
        t3: {"id": hod_id, "no": spread(base_cells[hod_id], 2), "max": 3, "first": False, "last": True},
    }
    for who, p in plan.items():
        p["star"] = free_cells(base_cells[p["id"]] + p["no"], 3)
        check(len(p["no"]) >= 2, f"{who}: {len(p['no'])} slots they REALLY had in the baseline will be marked 'not available'", str(p["no"]))

    # ------------------------------------------------------------------------------------------------ 2) the window opens
    log("\n== 2) the HOD opens the preference window ==")
    people = await db.fetch("select id from public.profiles where school_id = $1 and is_active and role in ('teacher','hod')", school)
    total = len(people)
    async with Acting(HOD_NAME):
        pre = (await client.get("/api/prefs/me")).json()
        check(pre["can_edit"] is False and "No preference window" in pre["reason"], "before any window nobody can edit (the HOD as a teacher included)")
        deadline = (datetime.now(vw.IST) + timedelta(days=3)).strftime("%Y-%m-%dT%H:%M")
        opened = await client.post("/api/prefs/window/open", json={"deadline": deadline, "note": "Phase 6 self-test"})
    check(opened.status_code == 201, "the HOD opens a window", opened.text[:100])
    wid = uuid.UUID(opened.json()["window"]["id"])
    check(opened.json()["notified"] == total - 1, f"all {total - 1} other teachers of the school were notified (not the HOD, not the admin)", str(opened.json()["notified"]))
    check(await notes(hod_id, "pref_window_open") == 0 and await notes(uuid.UUID(admin_id), "pref_window_open") == 0, "the HOD who opened it and the admin were NOT spammed")
    again = await as_(HOD_NAME, "POST", "/api/prefs/window/open", json={"deadline": deadline})
    check(again.status_code == 409 and again.json()["code"] == "ALREADY_OPEN", "a second window cannot be opened")

    # ------------------------------------------------------------------------------------------------ 3) teachers fill their grids
    log("\n== 3) three teachers fill their preferences (their own sessions) ==")
    for who, p in plan.items():
        body = {
            "slots": [{"day": d, "period": pn, "kind": "unavailable"} for d, pn in p["no"]] + [{"day": d, "period": pn, "kind": "preferred"} for d, pn in p["star"]],
            "prefer_first": p["first"], "avoid_last": p["last"], "max_per_day": p["max"],
        }
        r = await as_(who, "PUT", "/api/prefs/me", json=body)
        check(r.status_code == 200 and r.json()["submitted"], f"{who} saved: {r.json().get('summary') if r.status_code == 200 else r.text[:80]}")
        p["body"] = body
    over9 = [{"day": d, "period": pn, "kind": "unavailable"} for d in (1, 2, 3) for pn in (1, 2, 3)]
    refused = await as_(t2, "PUT", "/api/prefs/me", json={"slots": over9})
    check(refused.status_code == 422 and refused.json()["code"] == "TOO_MANY_UNAVAILABLE" and "at most 8" in refused.json()["message"], "9 'not available' marks are refused with a clear message", refused.json().get("message", "")[:90])
    mine = (await as_(t2, "GET", "/api/prefs/me")).json()
    check(len(mine["slots"]) == len(plan[t2]["body"]["slots"]), "…and nothing of the refused save was kept")
    visit = await db.fetchrow("select full_name, visiting_days from public.profiles where school_id = $1 and is_visiting order by full_name limit 1", school)
    off_day = next(d for d in range(1, 7) if d not in visit["visiting_days"])
    bad = await as_(visit["full_name"], "PUT", "/api/prefs/me", json={"slots": [{"day": off_day, "period": 1, "kind": "unavailable"}]})
    check(bad.status_code == 422 and bad.json()["code"] == "VISITING_DAY", f"visiting faculty ({visit['full_name']}): a slot on a day they do not come is refused - the rule stays separate", bad.json().get("message", "")[:80])
    okv = await as_(visit["full_name"], "PUT", "/api/prefs/me", json={"slots": [{"day": visit["visiting_days"][0], "period": 1, "kind": "preferred"}]})
    check(okv.status_code == 200 and okv.json()["locked_days"] == [d for d in range(1, 7) if d not in visit["visiting_days"] and d <= okv.json()["days"]], "…and on a visiting day it is saved")

    # the HOD's live list
    ov = (await as_(HOD_NAME, "GET", "/api/prefs/overview")).json()
    check(ov["filled"] == 4 and ov["total"] == total, f"live list: 4 of {total} filled", f"{ov['filled']}/{ov['total']}")
    check(all(t["filled"] for t in ov["teachers"] if t["name"] in (t1, t2, t3)) and sum(1 for t in ov["teachers"] if not t["filled"]) == total - 4, "…the other teachers are listed as 'not filled'")
    view = (await as_(HOD_NAME, "GET", f"/api/prefs/teacher/{t1_id}")).json()
    check(view["can_edit"] is False and len(view["slots"]) == len(plan[t1]["body"]["slots"]), "the HOD can open any teacher's preferences, read-only")
    other = await as_(t2, "GET", f"/api/prefs/teacher/{t1_id}")
    check(other.status_code == 403, "a teacher cannot read another teacher's preferences (403)")
    check((await as_(t2, "GET", "/api/prefs/overview")).status_code == 403, "a teacher cannot open the HOD's console (403)")

    # ------------------------------------------------------------------------------------------------ 3b) reminder, deadline, reopen
    log("\n== 3b) reminder (once, only the unfilled), auto-close at the deadline, reopen, extend, close ==")
    await db.execute("update public.preference_windows set opened_at = now() - interval '2 hours', deadline = now() + interval '10 minutes', reminder_sent_at = null where id = $1", wid)
    first = await prefflow.sweep(datetime.now(timezone.utc))
    second = await prefflow.sweep(datetime.now(timezone.utc))
    await notify.drain()
    check(first["reminded"] == total - 4 and second["reminded"] == 0, f"the reminder went to the {total - 4} teachers who had NOT filled - once", str(first))
    check(await notes(t1_id, "pref_reminder") == 0 and await notes(hod_id, "pref_reminder") == 0, "…and not to those who had filled")
    await db.execute("update public.preference_windows set deadline = now() - interval '1 minute' where id = $1", wid)
    end = await prefflow.sweep(datetime.now(timezone.utc))
    await notify.drain()
    row = await db.fetchrow("select status, closed_reason from public.preference_windows where id = $1", wid)
    check(end["closed"] == 1 and row["status"] == "closed" and row["closed_reason"] == "deadline", "at the deadline the window closes by itself")
    note = await db.fetchval("select body from public.notifications where user_id = $1 and kind = 'pref_window_closed' order by created_at desc limit 1", hod_id)
    check(f"4 of {total} teachers filled" in (note or ""), "…and the HOD is told how many filled", note or "")
    late = await as_(t1, "PUT", "/api/prefs/me", json=plan[t1]["body"])
    check(late.status_code == 409 and late.json()["code"] == "WINDOW_CLOSED", "after the deadline a teacher can view but not edit (409)")
    check((await as_(t1, "GET", "/api/prefs/me")).json()["can_edit"] is False, "…the screen says so")
    re = await as_(HOD_NAME, "POST", "/api/prefs/window/reopen", json={"deadline": (datetime.now(vw.IST) + timedelta(days=2)).strftime("%Y-%m-%dT%H:%M")})
    check(re.status_code == 200 and re.json()["notified"] == total - 4, f"reopened: only the {total - 4} unfilled teachers are told", str(re.json().get("notified")))
    ext = await as_(HOD_NAME, "POST", "/api/prefs/window/extend", json={"deadline": (datetime.now(vw.IST) + timedelta(days=5)).strftime("%Y-%m-%dT%H:%M")})
    check(ext.status_code == 200 and ext.json()["notified"] == total - 4, "the deadline can be extended (the unfilled are told)")
    cl = await as_(HOD_NAME, "POST", "/api/prefs/window/close", json={})
    check(cl.status_code == 200 and cl.json()["window"]["closed_reason"] == "manual", "the HOD can close it by hand")

    # ------------------------------------------------------------------------------------------------ 4) generation with the preferences
    log("\n== 4) generation WITH the preferences ==")
    pre_run = await api.post("/api/generation/runs", {"school_id": sid, "name": "P6 precheck", "time_limit_s": 120})
    check(not [f for f in pre_run["precheck"] if f["kind"] in ("teacher_preferences", "no_cell")], "no preference makes the timetable impossible (the precheck is quiet)")
    done = await wait_run(api, pre_run["id"])
    best = done["solutions"][0]
    cells = await teacher_cells(best["id"], school)
    classes = best["breakdown"]["classes"]
    check(all(v["percent"] == 100 and v["missing"] == 0 for v in classes.values()), "every lecture of all 8 classes is still placed (100 %)", str({v["label"]: v["percent"] for v in classes.values() if v["percent"] < 100}))
    for who, p in plan.items():
        got = set(cells[p["id"]])
        check(not (got & set(p["no"])), f"{who}: NOT ONE lecture on a 'not available' slot ({len(p['no'])} marked, {len(got)} periods taught)")
        per_day = Counter(d for d, _p in got)
        if p["max"]:
            check(max(per_day.values()) <= p["max"], f"{who}: never more than {p['max']} lectures in a day", str(dict(per_day)))
    check(not best["breakdown"]["unplaced"], "nothing left unplaced")
    await api.post(f"/api/generation/solutions/{best['id']}/use")
    rep = await api.get(f"/api/conflicts{q}")
    check(rep["counts"]["clashes"] == 0, "after using it: 0 clashes in the conflict report (so no 'not available' / daily-maximum clash either)")
    rule = next(r for r in best["breakdown"]["rules"] if r["rule"] == "preferred_slots")
    names = {t["name"]: t for t in rule["teachers"]}
    check(rule["weight"] == 60 and rule["detail"] and {t1, t2, t3} <= set(names), "the breakdown shows the new rule 'Teachers' preferred slots' with the detail and per-teacher result", rule["detail"])
    stars_on = sum(t["stars"][0] for t in rule["teachers"] if t["stars"])
    possible = sum(t["stars"][1] for t in rule["teachers"] if t["stars"])
    log(f"   stars honoured with the default weight: {stars_on} of {possible}; satisfaction {rule['satisfaction']}")

    off = {r["rule"]: 0 for r in best["breakdown"]["rules"]}
    off_run = await run_generation("P6 stars off", {**off, "preferred_slots": 0, "class_gaps": 80, "twice_a_day": 70})
    off_best = off_run["solutions"][0]
    off_rule = next(r for r in off_best["breakdown"]["rules"] if r["rule"] == "preferred_slots")
    stars_off = sum(t["stars"][0] for t in off_rule["teachers"] if t["stars"])
    check(stars_on >= stars_off, f"the weight steers the stars: honoured {stars_on} with weight 60 vs {stars_off} with weight 0")
    off_cells = await teacher_cells(off_best["id"], school)
    check(all(not (set(off_cells[p["id"]]) & set(p["no"])) for p in plan.values()), "even with the star weight at 0 a 'not available' slot is NEVER used (it is a hard rule, not a weight)")

    # ------------------------------------------------------------------------------------------------ 5) the manual grid
    log("\n== 5) the manual grid refuses a 'not available' cell and a crossed daily maximum ==")
    listing = await api.get(f"/api/drafts?school_id={sid}")
    primaries = {c["label"]: next(d for d in c["drafts"] if d["is_primary"]) for c in listing["classes"]}
    target = None
    for label, d in primaries.items():
        det = await api.get(f"/api/drafts/{d['id']}")
        for pl in det["placements"]:
            if pl["is_locked"] or pl["length"] != 1 or pl["parallel"]:
                continue
            if any(e["teacher_id"] == str(t1_id) for e in pl["entries"]):
                target = (d["id"], pl)
                break
        if target:
            break
    check(target is not None, f"found one of {t1}'s single lectures in a primary draft")
    draft_id, pl = target
    no_day, no_period = plan[t1]["no"][0]
    opts = await api.post(f"/api/drafts/{draft_id}/options", {"placement_key": pl["key"]})
    cell = opts["cells"][f"{no_day}:{no_period}"]
    check(cell["state"] == "blocked" and any("not available" in m for m in cell["hard"]), "the grid colours a 'not available' cell RED with the reason", (cell["hard"] or [""])[0][:90])
    mv = await client.patch(f"/api/drafts/{draft_id}/placements/{pl['key']}", json={"day": no_day, "start_period": no_period})
    check(mv.status_code == 422 and "marked" in mv.json()["message"] and "not available" in mv.json()["message"], "dropping the lecture there is refused (422)", mv.json()["message"][:100])
    still = next(p for p in (await api.get(f"/api/drafts/{draft_id}"))["placements"] if p["key"] == pl["key"])
    check((still["day"], still["start_period"]) == (pl["day"], pl["start_period"]), "…and the lecture stayed where it was")
    # the HOD-teacher's daily maximum
    t3_days = Counter(d for d, _p in cells[hod_id])
    full_day = next((d for d, n in t3_days.items() if n >= 3), None)
    seen_max = False
    if full_day is not None:
        for label, d in primaries.items():
            det = await api.get(f"/api/drafts/{d['id']}")
            mine = [p for p in det["placements"] if not p["is_locked"] and p["length"] == 1 and not p["parallel"] and any(e["teacher_id"] == str(hod_id) for e in p["entries"]) and p["day"] != full_day]
            if mine:
                o = await api.post(f"/api/drafts/{d['id']}/options", {"placement_key": mine[0]["key"]})
                seen_max = any(any("asked for at most 3 lecture(s) a day" in m for m in c["hard"]) for k, c in o["cells"].items() if k.startswith(f"{full_day}:"))
                break
    check(seen_max or full_day is None, f"{t3}'s daily maximum of 3 blocks a 4th lecture on their full day", f"day {full_day}")

    # ------------------------------------------------------------------------------------------------ 6) publish + reports
    log("\n== 6) publish, then every report against independent arithmetic ==")
    pub = await api.post("/api/publish", {"scope": "all", "school_id": sid})
    check(len(pub["published"]) == 8, "all 8 classes published", str(len(pub["published"])))
    wl = await api.get(f"/api/reports/workload{q}")
    prow = await db.fetch("select teacher_id, co_teacher_id, group_key, id, length, day from public.published_entries where school_id = $1", school)
    expect: dict[uuid.UUID, int] = defaultdict(int)
    seen: set[tuple[uuid.UUID, uuid.UUID]] = set()
    for r in prow:
        key = r["group_key"] or r["id"]
        for t in (r["teacher_id"], r["co_teacher_id"]):
            if t is not None and (t, key) not in seen:
                seen.add((t, key))
                expect[t] += r["length"]
    mism = [(r["teacher"], r["periods"], expect.get(uuid.UUID(r["teacher_id"]), 0)) for r in wl["rows"] if r["periods"] != expect.get(uuid.UUID(r["teacher_id"]), 0)]
    check(not mism and len(wl["rows"]) == total, f"workload: all {total} teachers' weekly periods equal an independent count of the published rows", str(mism[:3]))
    check(wl["summary"]["norm_hours"] == 16 and wl["summary"]["norm_source"] == "default" and wl["norm"]["min_hours"] == 12, "workload is judged against the default 16 h norm (configurable)")
    statuses = {r["status"] for r in wl["rows"]}
    check(all((r["status"] == "visiting") if r["is_visiting"] else (r["status"] == ("over" if r["hours"] > 16 else "under" if r["hours"] < 12 else "within")) for r in wl["rows"]), "every status follows the norm bands", str(statuses))
    top = max(r["hours"] for r in wl["rows"])
    rk = await api.get(f"/api/reports/rankings{q}&limit=5")
    check(rk["lists"]["busiest"][0]["hours"] == top and len(rk["lists"]["busiest"]) == 5, f"busiest teacher = {rk['lists']['busiest'][0]['teacher']} with {top} h; the list holds 5")
    most_free = max(r["free_periods"] for r in wl["rows"])
    check(rk["lists"]["most_free"][0]["free_periods"] == most_free, f"most free = {rk['lists']['most_free'][0]['teacher']} ({most_free} free periods)")

    # norm change -> numbers move
    await api.put("/api/reports/norms", {"scope": "college", "norm_hours": 8, "min_hours": 4, "label": "Self-test norm"})
    wl8 = await api.get(f"/api/reports/workload{q}")
    check(wl8["summary"]["norm_hours"] == 8 and wl8["summary"]["norm_source"] == "college" and wl8["summary"]["over"] >= wl["summary"]["over"], "changing the norm changes every status at once", f"over: {wl['summary']['over']} -> {wl8['summary']['over']}")
    await db.execute("delete from public.workload_norms where school_id is null")

    # free-teacher finder against the published rows
    day = nxt(1)
    wd = day.isoweekday()
    published = await db.fetch("select teacher_id, co_teacher_id, start_period, length from public.published_entries where school_id = $1 and day = $2", school, wd)
    busy_at: dict[int, set[uuid.UUID]] = defaultdict(set)
    for r in published:
        for t in (r["teacher_id"], r["co_teacher_id"]):
            if t is not None:
                for p in range(r["start_period"], r["start_period"] + r["length"]):
                    busy_at[p].add(t)
    visiting_off = {r["id"] for r in await db.fetch("select id, visiting_days from public.profiles where school_id = $1 and is_visiting", school) if wd not in r["visiting_days"]}
    period = max(range(1, 7), key=lambda p: len(busy_at[p]))
    free_expected = {r["id"] for r in people} - busy_at[period] - visiting_off
    ft = await api.get(f"/api/reports/free-teachers{q}&date={day.isoformat()}&period={period}")
    got_ids = {uuid.UUID(t["teacher_id"]) for t in ft["rows"]}
    check(got_ids == free_expected, f"free-teacher finder, {day:%a %d %b} period {period}: exactly the {len(free_expected)} teachers an independent check finds free", f"{len(got_ids)} listed")
    check(all(t["free_text"] and len(t["free_mask"]) == 6 for t in ft["rows"]) and ft["schools"][0]["time"], "…each with name, specialization, the periods they are free, with timings")
    nowr = await api.get(f"/api/reports/free-teachers{q}&now=true")
    check(nowr["now"] is True and isinstance(nowr["rows"], list), "…and 'right now' works too", "; ".join(nowr["notes"])[:90])

    # absences / substitutions / cancels -> the monthly reports
    busy_teacher = max(expect, key=expect.get)
    abs_day = Counter(r["day"] for r in prow if busy_teacher in (r["teacher_id"], r["co_teacher_id"])).most_common(1)[0][0]
    mday = nxt(abs_day)
    ab = await api.post("/api/absence", {"teacher_id": str(busy_teacher), "date": mday.isoformat(), "reason": "P6 self-test"})
    first_open = next(a for a in ab["affected"] if a["status"] == "open")
    cands = (await api.get(f"/api/subs/{first_open['id']}/candidates"))["candidates"]
    await api.post(f"/api/subs/{first_open['id']}/assign", {"teacher_id": cands[0]["teacher_id"]})
    month = mday.strftime("%Y-%m")
    sb = await api.get(f"/api/reports/substitutions{q}&month={month}")
    check(sb["summary"]["absences"] == 1 and sb["summary"]["needs"] == ab["to_cover"] and sb["summary"]["filled"] == 1, f"substitutions report: 1 absence, {ab['to_cover']} lecture(s) to cover, 1 filled", str({k: sb['summary'][k] for k in ('absences', 'needs', 'filled', 'fill_rate')}))
    check(sb["summary"]["fill_rate"] == round(100 / ab["to_cover"]) and any(r["subs_covered"] == 1 for r in sb["rows"]), "…the fill rate and the substitute's count are right")
    tue = nxt(2)
    cls_id, day_view = None, {"adjustable_starts": []}
    for c in (await api.get(f"/api/drafts?school_id={sid}"))["classes"]:
        dv = await api.get(f"/api/adjust?semester_id={c['semester_id']}&date={tue.isoformat()}")
        if dv["adjustable_starts"]:
            cls_id, day_view = c["semester_id"], dv
            break
    check(cls_id is not None, "a class with a lecture to cancel was found on the chosen Tuesday")
    if cls_id:
        before_h = await api.get(f"/api/reports/hours{q}&class_id={cls_id}")
        lost_before = sum(r["lost"] for r in before_h["rows"])
        await api.post("/api/adjust/cancel", {"semester_id": cls_id, "date": tue.isoformat(), "start_period": day_view["adjustable_starts"][0], "reason": "P6 self-test"})
        cx = await api.get(f"/api/reports/cancelled{q}&month={tue.strftime('%Y-%m')}")
        check(cx["summary"]["total"] == 1 and cx["rows"][0]["reason"] == "P6 self-test", "cancelled report: the one-day cancel is listed with its reason")
        after_h = await api.get(f"/api/reports/hours{q}&class_id={cls_id}")
        check(sum(r["lost"] for r in after_h["rows"]) == lost_before + 1, "…and counts as a lost hour in the subject-hours report")
    hours_rep = await api.get(f"/api/reports/hours{q}")
    phase4 = await api.get(f"/api/hours/subjects{q}")
    p4 = {s["subject_id"]: s for c in phase4["classes"] for s in c["subjects"]}
    diff = [(r["subject"], r["required"], p4[next(k for k, v in p4.items() if v["name"] == r["subject"] and v["class_label"] == r["class"])]["required"]) for r in hours_rep["rows"] if r["status"] in ("ok", "short") and any(
        v["name"] == r["subject"] and v["class_label"] == r["class"] and (v["required"], v["delivered"], v["deliverable"], v["shortfall"]) != (r["required"], r["delivered"], r["deliverable"], r["shortfall"]) for v in p4.values())]
    check(not diff and hours_rep["summary"]["short_hours"] == phase4["summary"]["short_hours"], f"subject-hours report equals the Semester Hours screen for all {len(hours_rep['rows'])} subjects", str(diff[:2]))
    cmp_ = await api.get(f"/api/reports/compare?month={month}")
    mine = next(r for r in cmp_["rows"] if r["school_id"] == sid)
    check(mine["teachers"] == total and mine["classes"] == 8 and mine["published_classes"] == 8 and mine["absences"] == 1, "school comparison row of School of IT matches (17 teachers, 8 classes, 8 published, 1 absence)")
    check(len(cmp_["rows"]) == len(await db.fetch("select 1 from public.schools where is_active")), "…the admin comparison has one row per active school")

    # the Excel files
    for name, extra in (("workload", ""), ("rankings", "&limit=5"), ("free-teachers", f"&date={day.isoformat()}&period={period}"), ("hours", ""),
                        ("cancelled", f"&month={tue.strftime('%Y-%m')}"), ("substitutions", f"&month={month}"), ("compare", f"&month={month}")):
        rj = await api.get(f"/api/reports/{name}{q}{extra}")
        r = await client.get(f"/api/reports/export/{name}{q}{extra}")
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.worksheets[0]
        n = 0
        for row in range(1, 40):
            if ws.cell(row, 1).value == rj["columns"][0]["label"]:
                n = sum(1 for rr in range(row + 1, ws.max_row + 1) if ws.cell(rr, 1).value not in (None, ""))
                break
        check(r.status_code == 200 and n == len(rj["rows"]), f"Excel '{name}': a real workbook with the same {len(rj['rows'])} rows as the screen")

    # ------------------------------------------------------------------------------------------------ 7) HOD locked to his school
    log("\n== 7) the HOD is locked to his own school in every report ==")
    stranger = str(uuid.uuid4())
    locked = 0
    for name in ("workload", "rankings", "free-teachers", "hours", "cancelled", "substitutions", "compare"):
        r1 = await as_(HOD_NAME, "GET", f"/api/reports/{name}", params={"school_id": stranger})
        r2 = await as_(HOD_NAME, "GET", f"/api/reports/export/{name}", params={"school_id": stranger})
        mine = (await as_(HOD_NAME, "GET", f"/api/reports/{name}")).json()
        locked += r1.status_code == 403 and r2.status_code == 403 and [s["id"] for s in mine["scope"]["schools"]] == [sid]
    check(locked == 7, "all 7 reports + their Excel files: another school -> 403, and without a filter the HOD gets only his own school")
    check((await as_(t1, "GET", "/api/reports/workload")).status_code == 403 and (await as_(t1, "GET", "/api/audit")).status_code == 403, "a teacher gets neither reports nor the audit log (403)")
    check((await as_(HOD_NAME, "GET", "/api/audit")).status_code == 403, "the HOD cannot read the audit log (admin only)")

    # ------------------------------------------------------------------------------------------------ 8) audit
    log("\n== 8) the audit screen shows real history ==")
    fac = await api.get("/api/audit/facets")
    check(fac["total"] > 20, f"the log holds {fac['total']} entries in all")
    oldest = await api.get("/api/audit", params={"page_size": 5, "page": (await api.get("/api/audit", params={"page_size": 5}))["pages"]})
    check(oldest["entries"] and all(e["at"] < started.astimezone(vw.IST).isoformat() for e in oldest["entries"]), "the oldest page is real history from BEFORE this test (the seed, the earlier phases)", oldest["entries"][-1]["summary"][:70])
    check(all(e["actor"]["email"] for e in oldest["entries"]), "…every entry says who did it")
    mine_a = await api.get("/api/audit", params={"actor_id": hod_id, "entity": "preference_window"})
    check({e["action"] for e in mine_a["entries"]} >= {"open", "reopen", "extend", "close"}, "this test's window actions appear under the HOD", str(sorted({e['action'] for e in mine_a['entries']})))
    mine_p = await api.get("/api/audit", params={"actor_id": str(t1_id), "entity": "preference"})
    check(mine_p["total"] >= 1 and "saved their preferences" in mine_p["entries"][0]["summary"], f"{t1}'s save appears, with details", mine_p["entries"][0]["summary"][:80])
    pg1 = await api.get("/api/audit", params={"page_size": 10})
    pg2 = await api.get("/api/audit", params={"page_size": 10, "page": 2})
    check(len(pg1["entries"]) == 10 and not ({e["id"] for e in pg1["entries"]} & {e["id"] for e in pg2["entries"]}) and pg1["pages"] > 5, f"pagination: {pg1['pages']} pages of 10, no overlap")
    rng = await api.get("/api/audit", params={"date_from": started.astimezone(vw.IST).date().isoformat(), "date_to": started.astimezone(vw.IST).date().isoformat(), "page_size": 100})
    check(rng["total"] >= 20 and all(e["at"][:10] == started.astimezone(vw.IST).date().isoformat() for e in rng["entries"]), f"date filter: {rng['total']} entries today, all dated today")
    import time as _t

    t0 = _t.monotonic()
    await api.get("/api/audit", params={"q": "published", "page_size": 50})
    check(_t.monotonic() - t0 < 3.0, "a text search over the whole log answers in under 3 seconds", f"{_t.monotonic() - t0:.2f}s")


# ====================================================================================================
async def cleanup(api: Api, started, original: dict[str, int | None], before: dict) -> None:
    log("\n== cleanup: every test artifact goes away, the preferences are restored ==")
    tree = await api.get("/api/structure/tree")
    school = uuid.UUID(next(s for s in tree["schools"] if s["name"] == "School of IT")["id"])
    sid = str(school)
    for sql in (
        "delete from public.substitution_offers where school_id = $1", "delete from public.substitutions where school_id = $1",
        "delete from public.day_adjustments where school_id = $1", "delete from public.teacher_absences where school_id = $1",
        "delete from public.leave_requests where school_id = $1", "delete from public.extra_lectures where school_id = $1",
        "delete from public.publish_events where school_id = $1", "delete from public.publications where school_id = $1",
        "delete from public.preference_windows where school_id = $1", "delete from public.generation_settings where school_id = $1",
    ):
        await db.execute(sql, school)
    await db.execute("delete from public.notifications where created_at >= $1", started)
    listing = await api.get(f"/api/drafts?school_id={sid}")
    for c in listing["classes"]:
        for d in c["drafts"]:
            await api.call("DELETE", f"/api/drafts/{d['id']}")
    runs = await api.get(f"/api/generation/runs?school_id={sid}")
    for r in runs["runs"]:
        if r["status"] == "running":
            await api.post(f"/api/generation/runs/{r['id']}/cancel")
            await wait_run(api, r["id"])
        await api.call("DELETE", f"/api/generation/runs/{r['id']}")
    restored = 0
    for sid_, weekly in original.items():
        n = await db.execute("update public.subjects set weekly_lectures = $2 where id = $1 and weekly_lectures is distinct from $2", uuid.UUID(sid_), weekly)
        restored += int(n.split()[-1])
    # the preferences go back to EXACTLY what they were before the test
    await db.execute("delete from public.teacher_pref_slots")
    await db.execute("delete from public.teacher_preferences")
    await db.execute("delete from public.preference_settings")
    await db.execute("delete from public.workload_norms")
    for table, rows in (("teacher_preferences", before["prefs"]), ("teacher_pref_slots", before["slots"]), ("preference_settings", before["settings"]), ("workload_norms", before["norms"])):
        for r in rows:
            cols = list(r)
            await db.execute(f"insert into public.{table} ({', '.join(cols)}) values ({', '.join('$' + str(i + 1) for i in range(len(cols)))})", *[r[c] for c in cols])
    await db.execute("delete from public.audit_log where at >= $1", started)
    left = await db.fetchrow(
        """
        select (select count(*) from public.timetable_drafts) d, (select count(*) from public.generation_runs) g, (select count(*) from public.publications) p,
               (select count(*) from public.published_entries) pe, (select count(*) from public.notifications) n, (select count(*) from public.substitutions) s,
               (select count(*) from public.teacher_absences) a, (select count(*) from public.day_adjustments) j,
               (select count(*) from public.subjects where weekly_lectures is not null) w,
               (select count(*) from public.preference_windows) pw, (select count(*) from public.preference_responses) pr
        """
    )
    check(all(v == 0 for v in dict(left).values()), f"drafts, runs, publications, notifications, absences, substitutions, windows and {restored} weekly-lectures overrides removed", str(dict(left)))
    now = await snapshot()
    check(now["prefs"] == before["prefs"] and now["slots"] == before["slots"] and now["settings"] == before["settings"] and now["norms"] == before["norms"],
          f"the preferences, limit and norm tables are exactly as before ({len(before['prefs'])} preference rows then, {len(now['prefs'])} now)")


if __name__ == "__main__":
    sys.exit(asyncio.run(main()))
